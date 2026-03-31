import openpyxl
book = openpyxl.load_workbook("C:\\Users\VanshAhuja\Downloads\download (18).xlsx")
sheet = book.active

for i in range(1,sheet.max_row+1):
    for j in range(1,sheet.max_column+1):
        if sheet.cell(row=i,column=j).value == "Apple":
            print(i,j)
            price =sheet.cell(row=i,column=j+2).value
            print(price)
            updated_price = sheet.cell(row=i, column=j + 2).value = 1000
            print(updated_price)
            assert updated_price == 1000

book.save("C:\\Users\VanshAhuja\Downloads\download (18).xlsx")
