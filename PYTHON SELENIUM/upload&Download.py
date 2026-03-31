from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions
from selenium.webdriver.support.wait import WebDriverWait

chrome_options = webdriver.ChromeOptions()
chrome_options.add_argument('headless')

driver = webdriver.Chrome(options=chrome_options)
driver.implicitly_wait(5)

file_path = r"C:\\Users\VanshAhuja\Downloads\download (18).xlsx"

driver.get("https://rahulshettyacademy.com/upload-download-test/")
driver.maximize_window()
fruit_name = "Apple"

driver.find_element(By.CSS_SELECTOR,"#downloadButton").click()
import openpyxl
book = openpyxl.load_workbook(file_path)
sheet = book.active
for i in range(1,sheet.max_row+1):
    for j in range(1,sheet.max_column+1):
        if sheet.cell(row=i,column=j).value == fruit_name:
            print(i,j)
            price =sheet.cell(row=i,column=j+2).value
            print(price)
            updated_price = sheet.cell(row=i, column=j + 2).value = 4000
            print(updated_price)
            assert updated_price == 4000
book.save(file_path)

file_input = driver.find_element(By.XPATH,"//input[@type='file']")
file_input.send_keys(file_path)
#why we have used send keys above because file upload button is not text, but we have used send keys
# and it will work if we send path of file where it is present in your system


#basically if we can give path of file which we want to upload in send keys method, So selenium
# takes path file and goes in system and uploads the file for you

wait = WebDriverWait(driver,10)

message = (By.XPATH,"//div[text()='Updated Excel Data Successfully.']")

wait.until(expected_conditions.visibility_of_element_located(message))

print(driver.find_element(By.XPATH,"//div[text()='Updated Excel Data Successfully.']").text)

print(driver.find_element(By.XPATH,"//div[text()='"+fruit_name+"']/parent::div/parent::div/div[4]").text)



