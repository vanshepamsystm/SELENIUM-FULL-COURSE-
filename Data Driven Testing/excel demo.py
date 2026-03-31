import openpyxl
book = openpyxl.load_workbook('C:\\Users\VanshAhuja\Desktop\Book1.xlsx') #loaded workbook
sheet = book.active
cell = sheet.cell(row=1,column=2)
print("Row 1 Column 2 Value: ",cell.value)     #first name

#to write something in sheet

sheet.cell(row=2,column=2).value = "Vansh"
print(sheet.cell(row=2,column=2).value)

print(sheet.max_row)     #5
print(sheet.max_column)     #3

print(sheet['A2'].value)


#below is script of printing value of sheet of 1st column all rows
for i in range(1,sheet.max_row+1):
    print(sheet.cell(row=i,column=1).value)

#this is to print all values
for i in range(1,sheet.max_row+1):
    for j in range(1,sheet.max_column+1):
        print(sheet.cell(row=i,column=j).value)



# this is how you print only 1 test case data
for i in range(1,sheet.max_row+1): #to get rows
    if sheet.cell(row=i,column=1).value == "testcase2":
        for j in range(2,sheet.max_column+1):   #to get columns
            print(sheet.cell(row=i,column=j).value)






















