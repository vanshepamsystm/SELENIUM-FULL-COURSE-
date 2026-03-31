import openpyxl
#have to tell where is file that is basic thing

book = openpyxl.load_workbook("C:\\Users\VanshAhuja\Desktop\Book3.xlsx")
# in above we have passed the path of the workbook using method load workbook

#few points
#-> now book object knows where is your workbook
#->out of that workbook you need to select a sheet which is active
#-> in Excel we can have multiple sheets
#-> to get control to active sheet we use method book.active

sheet = book.active

#-> in python each row and column treated started from number one(1)


# if you want to get cell below is method and you have to specify the row and column of cell

cell = sheet.cell(row=1,column=2)   #selected specific cell

print(cell.value) # printing value of the cell



#---------=================Moving onto how to write in the cell========================---------------

# this is how we assign value to a specific cell or enter data
sheet.cell(row=2,column=2).value = "Vansh"
print(sheet.cell(row=2,column=2).value )




#to know total number of rows and columns present in sheet

print(sheet.max_row)
print(sheet.max_column)



#shortcut way to print the value of cell

print(sheet['A5'].value)

#let now print all values present in Excel sheet using for loops

for i in range(1,sheet.max_row+1):
    for j in range(1,sheet.max_column+1):
        print(sheet.cell(row=i,column=j).value)

#printing data of only specific row with all columns
# we can use condtion here
print("testcase2 data only below")
for i in range(1,sheet.max_row+1):
    if sheet.cell(row=i,column=1).value == "Testcase2":
        for j in range(2,sheet.max_column+1):
            print(sheet.cell(row=i,column=j).value)
#this is how we print one test case data
print("Testcase2 data ended")

#----------------------==============how to store in dictionary===================----------------------
Dict = {}
for i in range(1,sheet.max_column+1):
    key = sheet.cell(row=1,column=i).value
    value = sheet.cell(row=3,column=i).value
    Dict[key] = value
print(Dict)
