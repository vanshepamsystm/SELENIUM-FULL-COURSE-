import openpyxl

book = openpyxl.load_workbook("C:\\Users\VanshAhuja\Desktop\Book2.xlsx")

sheet = book.active

print("Total Rows = ",sheet.max_row)
print("Total Columns = ",sheet.max_column)

print(sheet.cell(row=2,column=2).value)

sheet.cell(row=2,column=2).value = "Kriti"

print(sheet.cell(row=2,column=2).value)
#or
print(sheet['B2'].value)

#to get data only for tisha
for i in range(5,sheet.max_row+1):
    for j in range(1,sheet.max_column+1):
        print(sheet.cell(row=i,column=j).value)


for i in range(1,sheet.max_row+1):
    for j in range(1,sheet.max_column+1):
        print(sheet.cell(row=i,column=j).value)


print("Siddharth data: ")
for i in range(3,sheet.max_row-1):
    for j in range(2,sheet.max_column+1):
        print(sheet.cell(row=i,column=j).value)

Dict = {}
for i in range(3,sheet.max_row-1):
    for j in range(2,sheet.max_column+1):
        Dict[sheet.cell(row=1,column=j).value] =sheet.cell(row=i,column=j).value


print(Dict)